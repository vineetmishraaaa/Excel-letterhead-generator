import os
from flask import send_from_directory
import shutil
import subprocess
import tempfile
from io import BytesIO
from pathlib import Path
from threading import Lock
from urllib.parse import quote

import fitz
from flask import Flask, jsonify, render_template, request, send_file
from openpyxl import load_workbook
from openpyxl.styles.numbers import is_date_format
from werkzeug.utils import secure_filename


app = Flask(__name__)
@app.route("/google4b98075054d1ec81.html")
def google_verification():
    return send_from_directory("static", "google4b98075054d1ec81.html")
app.config["MAX_CONTENT_LENGTH"] = 50 * 1024 * 1024

BASE_DIR = Path(__file__).resolve().parent
UPLOAD_FOLDER = BASE_DIR / "uploads"
OUTPUT_FOLDER = BASE_DIR / "output"
UPLOAD_FOLDER.mkdir(exist_ok=True)
OUTPUT_FOLDER.mkdir(exist_ok=True)

ALLOWED_EXCEL = {".xlsx", ".xlsm"}
ALLOWED_LETTERHEAD = {".pdf"}

# One conversion at a time keeps the free 512 MB instance stable.
CONVERSION_LOCK = Lock()

LETTERHEAD_PAGE = 0
A4_WIDTH_CM = 21.0
A4_HEIGHT_CM = 29.7

# Tuned from the proven local workflow/reference PDF.
CONTENT_X_PT = 25.65
CONTENT_TOP_PT = 127.0
CONTENT_WIDTH_PT = 545.5
CONTENT_HEIGHT_PT = 345.0


# ============================================================
# WORKBOOK PREPARATION
# ============================================================

def prepare_workbook(src: Path, dst: Path) -> None:
    """Prepare a temporary workbook copy for Calc PDF export.

    The original upload is never changed.
    """
    keep_vba = src.suffix.lower() == ".xlsm"
    wb = load_workbook(src, keep_vba=keep_vba)

    for ws in wb.worksheets:
        if ws.sheet_state != "visible":
            continue

        setup = ws.page_setup

        # A4 portrait, matching the working desktop version.
        setup.paperSize = ws.PAPERSIZE_A4
        setup.orientation = ws.ORIENTATION_PORTRAIT

        # Critical: keep each worksheet on one printed page.
        setup.fitToWidth = 1
        setup.fitToHeight = 1
        ws.sheet_properties.pageSetUpPr.fitToPage = True

        # These only affect the temporary copy.
        setup.leftMargin = 0.20
        setup.rightMargin = 0.20
        setup.topMargin = 0.15
        setup.bottomMargin = 0.20
        ws.print_options.gridLines = False

        # Letterhead is supplied separately.
        setup.leftHeader = ""
        setup.centerHeader = ""
        setup.rightHeader = ""
        setup.leftFooter = ""
        setup.centerFooter = ""
        setup.rightFooter = ""

        # Preserve an existing print area; otherwise use used range.
        if not ws.print_area:
            ws.print_area = ws.calculate_dimension()

        # Excel/LibreOffice can display date formats differently.
        # Normalize real date-formatted cells to the expected format.
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is not None and is_date_format(cell.number_format):
                    cell.number_format = "dd-mm-yyyy"

    wb.save(dst)


def export_excel_to_pdf(src: Path, pdf_out: Path) -> None:
    """Convert the complete workbook with one-shot headless LibreOffice.

    This avoids the long-running UNO bridge that caused the original
    Render memory problem. The temporary workbook is deleted afterwards.
    """
    temp_dir = Path(tempfile.mkdtemp(prefix="excelpdf_"))
    prepared = temp_dir / "prepared.xlsx"
    profile = temp_dir / "lo_profile"

    try:
        prepare_workbook(src, prepared)
        profile.mkdir(exist_ok=True)

        profile_url = "file://" + quote(
            str(profile.resolve()).replace("\\", "/"),
            safe="/:"
        )

        env = os.environ.copy()
        env["HOME"] = "/tmp"
        env["SAL_USE_VCLPLUGIN"] = "svp"
        env["SAL_DISABLE_OPENCL"] = "1"
        env["MALLOC_TRIM_THRESHOLD_"] = "65536"

        cmd = [
            "libreoffice",
            "--headless",
            "--nologo",
            "--nodefault",
            "--nofirststartwizard",
            "--norestore",
            "--nolockcheck",
            f"-env:UserInstallation={profile_url}",
            "--convert-to",
            "pdf:calc_pdf_Export",
            "--outdir",
            str(temp_dir.resolve()),
            str(prepared.resolve()),
        ]

        result = subprocess.run(
            cmd,
            capture_output=True,
            text=True,
            timeout=180,
            env=env,
        )

        generated = prepared.with_suffix(".pdf")

        if result.returncode != 0 or not generated.exists():
            details = (result.stderr or result.stdout).strip()
            raise RuntimeError(
                details or "Excel to PDF conversion failed."
            )

        if pdf_out.exists():
            pdf_out.unlink()

        shutil.move(str(generated), str(pdf_out))

    finally:
        shutil.rmtree(temp_dir, ignore_errors=True)


# ============================================================
# PDF PLACEMENT
# ============================================================

def visible_content_bbox(page: fitz.Page) -> fitz.Rect:
    """Find Excel's visible vector content without rasterization."""
    rects = []

    for block in page.get_text("blocks"):
        if len(block) >= 4:
            rects.append(
                fitz.Rect(block[0], block[1], block[2], block[3])
            )

    for drawing in page.get_drawings():
        rect = drawing.get("rect")
        if rect:
            rects.append(rect)

    for image in page.get_images(full=True):
        try:
            rects.extend(page.get_image_rects(image[0]))
        except Exception:
            pass

    if not rects:
        return page.rect

    bbox = rects[0]
    for rect in rects[1:]:
        bbox |= rect

    padding = 3

    return fitz.Rect(
        max(page.rect.x0, bbox.x0 - padding),
        max(page.rect.y0, bbox.y0 - padding),
        min(page.rect.x1, bbox.x1 + padding),
        min(page.rect.y1, bbox.y1 + padding),
    )


def create_final_pdf(
    letterhead_pdf: Path,
    excel_pdf: Path,
    final_pdf: Path,
) -> int:
    """Apply the clean letterhead to every Excel PDF page."""
    letterhead_doc = fitz.open(str(letterhead_pdf))
    excel_doc = fitz.open(str(excel_pdf))
    output_doc = fitz.open()

    try:
        if len(letterhead_doc) <= LETTERHEAD_PAGE:
            raise ValueError(
                "Letterhead PDF must contain a first-page template."
            )

        if not excel_doc:
            raise ValueError(
                "Excel conversion produced no PDF pages."
            )

        for page_number, excel_page in enumerate(excel_doc):
            page = output_doc.new_page(
                width=A4_WIDTH_CM * 72 / 2.54,
                height=A4_HEIGHT_CM * 72 / 2.54,
            )

            # Full clean letterhead background.
            page.show_pdf_page(
                page.rect,
                letterhead_doc,
                LETTERHEAD_PAGE,
                keep_proportion=False,
            )

            # Find the worksheet content without creating a bitmap.
            crop = visible_content_bbox(excel_page)

            # Fixed body box is intentional. It matches the proven local
            # output much more closely than proportional scaling of Calc's
            # slightly different page geometry.
            destination = fitz.Rect(
                CONTENT_X_PT,
                CONTENT_TOP_PT,
                CONTENT_X_PT + CONTENT_WIDTH_PT,
                CONTENT_TOP_PT + CONTENT_HEIGHT_PT,
            )

            # We allow non-uniform scaling because the source renderer
            # (Calc) and desktop Excel produce slightly different page
            # geometry. The actual worksheet remains vector PDF content.
            page.show_pdf_page(
                destination,
                excel_doc,
                page_number,
                keep_proportion=False,
                clip=crop,
            )

        output_doc.save(
            str(final_pdf),
            garbage=4,
            deflate=True,
            clean=True,
        )

        return len(excel_doc)

    finally:
        output_doc.close()
        excel_doc.close()
        letterhead_doc.close()


# ============================================================
# WEB ROUTES
# ============================================================

@app.get("/")
def index():
    return render_template("index.html")


@app.get("/health")
def health():
    return jsonify({"status": "ok"})


@app.post("/upload")
def upload():
    excel_file = request.files.get("excel_file")
    letterhead_file = request.files.get("letterhead")

    if not excel_file or not letterhead_file:
        return jsonify({
            "error": "Please select both files."
        }), 400

    if not excel_file.filename or not letterhead_file.filename:
        return jsonify({
            "error": "Please select both files."
        }), 400

    excel_name = secure_filename(excel_file.filename)
    letterhead_name = secure_filename(letterhead_file.filename)

    if Path(excel_name).suffix.lower() not in ALLOWED_EXCEL:
        return jsonify({
            "error": "Please upload an .xlsx or .xlsm file."
        }), 400

    if Path(letterhead_name).suffix.lower() not in ALLOWED_LETTERHEAD:
        return jsonify({
            "error": "Please upload a PDF letterhead."
        }), 400

    # Serialize conversions to keep RAM predictable on free hosting.
    with CONVERSION_LOCK:
        with tempfile.TemporaryDirectory(
            prefix="request_"
        ) as request_dir:

            request_dir = Path(request_dir)

            excel_path = request_dir / excel_name
            letterhead_path = request_dir / letterhead_name
            excel_pdf = request_dir / "excel_export.pdf"
            final_pdf = request_dir / "final_letterhead.pdf"

            excel_file.save(excel_path)
            letterhead_file.save(letterhead_path)

            try:
                app.logger.info("Starting Excel conversion")

                export_excel_to_pdf(
                    excel_path,
                    excel_pdf,
                )

                app.logger.info("Starting letterhead merge")

                page_count = create_final_pdf(
                    letterhead_path,
                    excel_pdf,
                    final_pdf,
                )

                pdf_bytes = final_pdf.read_bytes()

                app.logger.info(
                    "Generated %s page(s)",
                    page_count,
                )

            except Exception as exc:
                app.logger.exception(
                    "PDF generation failed"
                )

                return jsonify({
                    "error": str(exc)
                }), 500

    download_name = (
        Path(excel_name).stem
        + "_WITH_LETTERHEAD.pdf"
    )

    return send_file(
        BytesIO(pdf_bytes),
        as_attachment=True,
        download_name=download_name,
        mimetype="application/pdf",
    )


if __name__ == "__main__":
    app.run(
        host="0.0.0.0",
        port=int(os.environ.get("PORT", "5000")),
        debug=False,
    )
